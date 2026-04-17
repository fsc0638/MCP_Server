"""
Employee Lookup Service — matches LINE users to company employee records.

Data source: workspace/department/同仁清單.xlsx
Columns: #(員編+姓名), 同仁姓名, 分機, 部門, 職稱, 信箱

Provides:
- Lookup by email or employee ID
- Department list extraction
- User context generation for onboarding
"""

import json
import logging
import os
import re
from pathlib import Path
from typing import Optional, Dict, Any, List

logger = logging.getLogger("MCP_Server.EmployeeLookup")

_EMPLOYEE_CACHE = None  # Lazy-loaded
_DEPARTMENTS_CACHE = None


def _get_xlsx_path() -> Path:
    """Resolve employee list path."""
    pr = os.getenv("PROJECT_ROOT", "")
    if pr:
        return Path(pr) / "workspace" / "department" / "同仁清單.xlsx"
    return Path(__file__).resolve().parents[2] / "workspace" / "department" / "同仁清單.xlsx"


def _open_xlsx_safely(xlsx_path: Path):
    """Open xlsx handling OneDrive/Excel file locks.

    On Windows, OneDrive sync or open Excel instance can produce PermissionError
    when we try to read. Workaround: copy the file to a temp location and open
    the copy. This is read-only from our perspective so it's safe.
    """
    import openpyxl
    import shutil
    import tempfile
    try:
        return openpyxl.load_workbook(str(xlsx_path), read_only=True, data_only=True)
    except PermissionError:
        logger.info(f"[EmployeeLookup] {xlsx_path.name} is locked, copying to temp for read")
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name
        try:
            shutil.copy2(str(xlsx_path), tmp_path)
            return openpyxl.load_workbook(tmp_path, read_only=True, data_only=True)
        finally:
            try:
                import os as _os
                _os.unlink(tmp_path)
            except Exception:
                pass


def _load_employees() -> List[Dict[str, Any]]:
    """Load and parse employee list from xlsx. Cached after first load."""
    global _EMPLOYEE_CACHE
    if _EMPLOYEE_CACHE is not None:
        return _EMPLOYEE_CACHE

    xlsx_path = _get_xlsx_path()
    if not xlsx_path.exists():
        logger.warning(f"[EmployeeLookup] File not found: {xlsx_path}")
        return []

    try:
        wb = _open_xlsx_safely(xlsx_path)
        ws = wb.active
        employees = []

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
            if not row or not row[1]:
                continue

            # Parse "員編　姓名" format (e.g., "0337　羅燕秋")
            raw_name = str(row[1]).strip()
            emp_id = ""
            name = raw_name
            match = re.match(r'^(\d+)\s+(.+)$', raw_name)
            if match:
                emp_id = match.group(1).strip()
                name = match.group(2).strip()

            # Parse department "代碼　名稱" format (e.g., "A100　稽核室")
            raw_dept = str(row[3]).strip() if row[3] else ""
            dept_code = ""
            dept_name = raw_dept
            dept_match = re.match(r'^([A-Z]\d+)\s+(.+)$', raw_dept)
            if dept_match:
                dept_code = dept_match.group(1).strip()
                dept_name = dept_match.group(2).strip()

            # Column 6 (index 6) = role (admin/editor/viewer), added by admin panel
            _role = str(row[6]).strip().lower() if len(row) > 6 and row[6] else ""

            employees.append({
                "employee_id": emp_id,
                "name": name,
                "extension": str(row[2]).strip() if row[2] else "",
                "department_code": dept_code,
                "department_name": dept_name,
                "department_full": raw_dept,
                "title": str(row[4]).strip() if row[4] else "",
                "email": str(row[5]).strip().lower() if row[5] else "",
                "role": _role,
            })

        wb.close()
        _EMPLOYEE_CACHE = employees
        logger.info(f"[EmployeeLookup] Loaded {len(employees)} employees")
        return employees

    except Exception as e:
        logger.error(f"[EmployeeLookup] Failed to load: {e}")
        return []


def save_employee_to_xlsx(employee_id: str, updates: Dict[str, Any]) -> bool:
    """Write updated employee data back to xlsx and refresh cache."""
    global _EMPLOYEE_CACHE
    xlsx_path = _get_xlsx_path()
    if not xlsx_path.exists():
        return False

    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(xlsx_path))
        ws = wb.active

        # Ensure header row has "角色" column (col 7 = G)
        header = [ws.cell(row=1, column=c).value for c in range(1, 8)]
        if len(header) < 7 or header[6] != "角色":
            ws.cell(row=1, column=7, value="角色")

        # Find the row by employee_id
        found_row = None
        for r in range(2, ws.max_row + 1):
            raw_name = str(ws.cell(row=r, column=2).value or "").strip()
            match = re.match(r'^(\d+)\s+', raw_name)
            if match and match.group(1).strip() == employee_id:
                found_row = r
                break

        if not found_row:
            wb.close()
            return False

        # Update cells
        if "name" in updates:
            ws.cell(row=found_row, column=2, value=f"{employee_id} {updates['name']}")
        if "extension" in updates:
            ws.cell(row=found_row, column=3, value=updates["extension"])
        if "department_code" in updates and "department_name" in updates:
            ws.cell(row=found_row, column=4, value=f"{updates['department_code']} {updates['department_name']}")
        elif "department_code" in updates:
            ws.cell(row=found_row, column=4, value=updates["department_code"])
        if "title" in updates:
            ws.cell(row=found_row, column=5, value=updates["title"])
        if "email" in updates:
            ws.cell(row=found_row, column=6, value=updates["email"])
        if "role" in updates:
            ws.cell(row=found_row, column=7, value=updates["role"])

        wb.save(str(xlsx_path))
        wb.close()

        # Invalidate cache so next load picks up changes
        _EMPLOYEE_CACHE = None
        logger.info(f"[EmployeeLookup] Updated employee {employee_id} in xlsx")
        return True

    except Exception as e:
        logger.error(f"[EmployeeLookup] Failed to save xlsx: {e}")
        return False


def get_departments() -> List[Dict[str, str]]:
    """Return sorted list of unique departments."""
    global _DEPARTMENTS_CACHE
    if _DEPARTMENTS_CACHE is not None:
        return _DEPARTMENTS_CACHE

    employees = _load_employees()
    seen = {}
    for emp in employees:
        code = emp["department_code"]
        if code and code not in seen:
            seen[code] = {
                "code": code,
                "name": emp["department_name"],
                "full": emp["department_full"],
            }
    _DEPARTMENTS_CACHE = sorted(seen.values(), key=lambda d: d["code"])
    return _DEPARTMENTS_CACHE


def lookup_by_email(email: str) -> Optional[Dict[str, Any]]:
    """Find employee by email address."""
    email = email.strip().lower()
    for emp in _load_employees():
        if emp["email"] == email:
            return emp
    return None


def lookup_by_employee_id(emp_id: str) -> Optional[Dict[str, Any]]:
    """Find employee by employee ID (numbers only)."""
    emp_id = emp_id.strip().lstrip("0")  # Normalize: "0337" → "337"
    for emp in _load_employees():
        stored_id = emp["employee_id"].lstrip("0")
        if stored_id == emp_id:
            return emp
    return None


def lookup_by_name(name: str) -> Optional[Dict[str, Any]]:
    """Find employee by exact name match."""
    name = name.strip()
    for emp in _load_employees():
        if emp["name"] == name:
            return emp
    return None


def lookup(query: str) -> Optional[Dict[str, Any]]:
    """Smart lookup: try email → employee ID → name."""
    query = query.strip()
    if not query:
        return None

    # Email pattern
    if "@" in query:
        return lookup_by_email(query)

    # Pure digits → employee ID
    if query.isdigit():
        return lookup_by_employee_id(query)

    # Otherwise try name
    return lookup_by_name(query)


def build_user_context(employee: Dict[str, Any], user_id: str, role: str = "editor") -> Dict[str, Any]:
    """Build user context JSON from matched employee record."""
    return {
        "user_id": user_id,
        "employee_id": employee.get("employee_id", ""),
        "name": employee.get("name", ""),
        "department": employee.get("department_full", ""),
        "department_code": employee.get("department_code", ""),
        "department_name": employee.get("department_name", ""),
        "title": employee.get("title", ""),
        "email": employee.get("email", ""),
        "extension": employee.get("extension", ""),
        "preferences": {
            "language": "繁體中文",
            "style": "適中",
            "primary_use": [],
        },
        "role": role,
        "groups": [],
        "skill_access": {
            "system": "all",
            "department": [employee.get("department_code", "")],
            "personal": True,
        },
        "onboarding_completed": True,
        "source": "line_bot",
    }


def save_user_context(user_context: Dict[str, Any]) -> Path:
    """Save user context to workspace/users/{user_id}.json."""
    pr = os.getenv("PROJECT_ROOT", "")
    if pr:
        users_dir = Path(pr) / "workspace" / "users"
    else:
        users_dir = Path(__file__).resolve().parents[2] / "workspace" / "users"
    users_dir.mkdir(parents=True, exist_ok=True)

    user_id = user_context.get("user_id", "unknown")
    path = users_dir / f"{user_id}.json"
    with open(path, "w", encoding="utf-8") as f:
        json.dump(user_context, f, ensure_ascii=False, indent=2)

    logger.info(f"[EmployeeLookup] User context saved: {path}")
    return path


def get_user_context(user_id: str) -> Optional[Dict[str, Any]]:
    """Load existing user context. Returns None if not found."""
    pr = os.getenv("PROJECT_ROOT", "")
    if pr:
        path = Path(pr) / "workspace" / "users" / f"{user_id}.json"
    else:
        path = Path(__file__).resolve().parents[2] / "workspace" / "users" / f"{user_id}.json"

    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None


def has_completed_onboarding(user_id: str) -> bool:
    """Check if user has completed onboarding."""
    ctx = get_user_context(user_id)
    return ctx is not None and ctx.get("onboarding_completed", False)
