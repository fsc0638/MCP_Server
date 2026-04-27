"""DOCX preview helpers with dependency-light fallbacks."""

from __future__ import annotations

import html
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Iterator
from xml.etree import ElementTree as ET

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}
W = f"{{{W_NS}}}"

PRINT_PAGE_HTML = """<!DOCTYPE html>
<html lang="zh-TW">
<head>
  <meta charset="UTF-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>{title}</title>
  <style>
    @page {{ size: A4; margin: 16mm 14mm; }}
    html, body {{
      margin: 0;
      padding: 0;
      color: #0f172a;
      background: #ffffff;
      font-family: "Segoe UI", "Microsoft JhengHei", "PingFang TC", "Noto Sans TC", sans-serif;
      font-size: 12pt;
      line-height: 1.7;
    }}
    body {{
      padding: 0;
    }}
    article {{
      width: 100%;
    }}
    h1, h2, h3, h4, h5, h6 {{
      line-height: 1.3;
      break-after: avoid-page;
    }}
    h1 {{ font-size: 22pt; }}
    h2 {{ font-size: 18pt; }}
    h3 {{ font-size: 15pt; }}
    p {{
      margin: 0 0 10pt;
      orphans: 2;
      widows: 2;
    }}
    ul, ol {{
      margin: 0 0 12pt 18pt;
      padding: 0;
    }}
    li + li {{
      margin-top: 4pt;
    }}
    .docx-preview-spacer {{
      height: 10pt;
    }}
    .docx-preview-title {{
      margin-top: 0;
      margin-bottom: 8pt;
    }}
    .docx-preview-subtitle {{
      color: #475569;
    }}
    .docx-preview-table-wrap {{
      margin: 12pt 0 14pt;
      overflow: hidden;
      border: 1px solid #dbe4f0;
      border-radius: 8pt;
    }}
    .docx-preview-table {{
      width: 100%;
      border-collapse: collapse;
      table-layout: fixed;
    }}
    .docx-preview-table th,
    .docx-preview-table td {{
      padding: 8pt 9pt;
      border-right: 1px solid #dbe4f0;
      border-bottom: 1px solid #dbe4f0;
      vertical-align: top;
      text-align: left;
      word-break: break-word;
    }}
    .docx-preview-table th:last-child,
    .docx-preview-table td:last-child {{
      border-right: none;
    }}
    .docx-preview-table thead th {{
      background: #eef4ff;
      font-weight: 700;
    }}
    .docx-preview-table tbody tr:last-child td {{
      border-bottom: none;
    }}
  </style>
</head>
<body>
  {body}
</body>
</html>
"""


def iter_block_items(parent: Any) -> Iterator[Any]:
    """Yield paragraphs and tables in document order for python-docx objects."""
    from docx.document import Document as DocumentType
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph

    if isinstance(parent, DocumentType):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise TypeError(f"Unsupported parent type: {type(parent)!r}")

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def render_docx_to_html(docx_path: str) -> str:
    """Convert a DOCX file into a safe HTML fragment for preview."""
    try:
        rendered = _render_docx_to_html_via_python_docx(docx_path)
        if rendered.strip():
            return rendered
    except Exception:
        pass

    blocks = _parse_docx_blocks(docx_path)
    rendered = _render_blocks_to_html(blocks)
    if rendered.strip():
        return rendered
    return '<article class="docx-preview-fragment"><p>No previewable content found.</p></article>'


def extract_docx_plain_text(docx_path: str) -> str:
    """Extract readable plain text from a DOCX file."""
    try:
        from docx import Document

        doc = Document(docx_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text and cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        text = "\n".join(paragraphs).strip()
        if text:
            return text
    except Exception:
        pass

    try:
        import docx2txt

        text = (docx2txt.process(docx_path) or "").strip()
        if text:
            return text
    except Exception:
        pass

    blocks = _parse_docx_blocks(docx_path)
    return _blocks_to_text(blocks).strip()


def convert_docx_to_pdf(docx_path: str, pdf_path: str) -> None:
    """Convert a DOCX file to PDF for inline preview."""
    source = Path(docx_path).resolve()
    target = Path(pdf_path).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    target.unlink(missing_ok=True)

    errors: list[str] = []
    converters = (
        _convert_via_docx2pdf,
        _convert_via_word_com,
        _convert_via_soffice,
        _convert_via_browser_print,
    )

    for converter in converters:
        try:
            converter(source, target)
            if target.exists():
                return
        except Exception as exc:
            errors.append(str(exc).strip() or converter.__name__)

    joined = " ; ".join(error for error in errors if error) or "No available DOCX-to-PDF converter."
    raise RuntimeError(joined)


def _render_docx_to_html_via_python_docx(docx_path: str) -> str:
    from docx import Document
    from docx.text.paragraph import Paragraph

    doc = Document(docx_path)
    parts: list[str] = []
    current_list_type: str | None = None

    def close_list() -> None:
        nonlocal current_list_type
        if current_list_type:
            parts.append(f"</{current_list_type}>")
            current_list_type = None

    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            text_html = _render_python_docx_paragraph_runs(block).strip()
            list_type = _resolve_python_docx_list_type(block)

            if list_type:
                if current_list_type != list_type:
                    close_list()
                    current_list_type = list_type
                    parts.append(f"<{list_type}>")
                if text_html:
                    parts.append(f"<li>{text_html}</li>")
                continue

            close_list()

            if not text_html:
                parts.append('<div class="docx-preview-spacer" aria-hidden="true"></div>')
                continue

            tag, css_class = _get_block_tag(block.style.name if block.style else "")
            if css_class:
                parts.append(f'<{tag} class="{css_class}">{text_html}</{tag}>')
            else:
                parts.append(f"<{tag}>{text_html}</{tag}>")
            continue

        close_list()
        table_html = _render_python_docx_table(block)
        if table_html:
            parts.append(table_html)

    close_list()

    if not parts:
        parts.append("<p>No previewable content found.</p>")

    return '<article class="docx-preview-fragment">' + "".join(parts) + "</article>"


def _resolve_python_docx_list_type(paragraph: Any) -> str | None:
    style_name = ((paragraph.style.name if paragraph.style else "") or "").strip().lower()
    p_pr = getattr(paragraph._p, "pPr", None)
    has_numbering = bool(p_pr is not None and getattr(p_pr, "numPr", None) is not None)

    if "list bullet" in style_name or "bullet" in style_name:
        return "ul"
    if "list number" in style_name or "number" in style_name:
        return "ol"
    if has_numbering:
        return "ol" if "number" in style_name else "ul"
    return None


def _render_python_docx_paragraph_runs(paragraph: Any) -> str:
    chunks: list[str] = []
    for run in paragraph.runs:
        text = _escape_run_text(run.text)
        if not text:
            continue
        if run.bold:
            text = f"<strong>{text}</strong>"
        if run.italic:
            text = f"<em>{text}</em>"
        if run.underline:
            text = f"<u>{text}</u>"
        chunks.append(text)

    if chunks:
        return "".join(chunks)
    return _escape_run_text(paragraph.text)


def _render_python_docx_table(table: Any) -> str:
    rows = list(table.rows)
    if not rows:
        return ""

    rendered_rows = [[_render_python_docx_table_cell(cell) for cell in row.cells] for row in rows]
    if not any(any(cell.strip() for cell in row) for row in rendered_rows):
        return ""

    head_cells = "".join(f"<th>{cell or '&nbsp;'}</th>" for cell in rendered_rows[0])
    body_rows = rendered_rows[1:] if len(rendered_rows) > 1 else []
    body_html = "".join(
        "<tr>" + "".join(f"<td>{cell or '&nbsp;'}</td>" for cell in row) + "</tr>"
        for row in body_rows
    )

    parts = [
        '<div class="docx-preview-table-wrap"><table class="docx-preview-table">',
        f"<thead><tr>{head_cells}</tr></thead>",
    ]
    if body_html:
        parts.append(f"<tbody>{body_html}</tbody>")
    parts.append("</table></div>")
    return "".join(parts)


def _render_python_docx_table_cell(cell: Any) -> str:
    from docx.text.paragraph import Paragraph

    chunks: list[str] = []
    for block in iter_block_items(cell):
        if isinstance(block, Paragraph):
            html_text = _render_python_docx_paragraph_runs(block).strip()
            if html_text:
                chunks.append(html_text)
        else:
            nested = _render_python_docx_table(block)
            if nested:
                chunks.append(nested)
    return "<br>".join(chunks)


def _parse_docx_blocks(docx_path: str) -> list[dict[str, Any]]:
    with zipfile.ZipFile(docx_path) as archive:
        document_root = _read_xml_part(archive, "word/document.xml")
        styles_root = _read_xml_part(archive, "word/styles.xml", required=False)
        numbering_root = _read_xml_part(archive, "word/numbering.xml", required=False)

    styles = _parse_styles(styles_root)
    numbering = _parse_numbering(numbering_root)

    body = document_root.find("w:body", NS)
    if body is None:
        return []

    blocks: list[dict[str, Any]] = []
    for child in body:
        block = _parse_block(child, styles, numbering)
        if block:
            blocks.append(block)
    return blocks


def _read_xml_part(archive: zipfile.ZipFile, name: str, required: bool = True) -> ET.Element | None:
    try:
        return ET.fromstring(archive.read(name))
    except KeyError:
        if required:
            raise
        return None


def _parse_styles(root: ET.Element | None) -> dict[str, str]:
    if root is None:
        return {}

    styles: dict[str, str] = {}
    for style in root.findall("w:style", NS):
        style_id = style.get(W + "styleId") or ""
        name_node = style.find("w:name", NS)
        style_name = name_node.get(W + "val") if name_node is not None else ""
        if style_id:
            styles[style_id] = style_name or style_id
    return styles


def _parse_numbering(root: ET.Element | None) -> dict[str, dict[str, str]]:
    if root is None:
        return {}

    abstract_formats: dict[str, dict[str, str]] = {}
    for abstract_num in root.findall("w:abstractNum", NS):
        abstract_id = abstract_num.get(W + "abstractNumId") or ""
        levels: dict[str, str] = {}
        for level in abstract_num.findall("w:lvl", NS):
            ilvl = level.get(W + "ilvl") or "0"
            fmt_node = level.find("w:numFmt", NS)
            num_fmt = fmt_node.get(W + "val") if fmt_node is not None else ""
            levels[ilvl] = num_fmt
        if abstract_id:
            abstract_formats[abstract_id] = levels

    numbering: dict[str, dict[str, str]] = {}
    for num in root.findall("w:num", NS):
        num_id = num.get(W + "numId") or ""
        abstract_ref = num.find("w:abstractNumId", NS)
        abstract_id = abstract_ref.get(W + "val") if abstract_ref is not None else ""
        if num_id and abstract_id in abstract_formats:
            numbering[num_id] = abstract_formats[abstract_id]
    return numbering


def _parse_block(element: ET.Element, styles: dict[str, str], numbering: dict[str, dict[str, str]]) -> dict[str, Any] | None:
    if element.tag == W + "p":
        return _parse_paragraph_block(element, styles, numbering)
    if element.tag == W + "tbl":
        return _parse_table_block(element, styles, numbering)
    return None


def _parse_paragraph_block(
    element: ET.Element,
    styles: dict[str, str],
    numbering: dict[str, dict[str, str]],
) -> dict[str, Any]:
    p_pr = element.find("w:pPr", NS)
    style_id = ""
    if p_pr is not None:
        style_node = p_pr.find("w:pStyle", NS)
        style_id = style_node.get(W + "val") if style_node is not None else ""
    style_name = styles.get(style_id, style_id)
    runs = _parse_runs_from_container(element)
    text = "".join(run["text"] for run in runs)
    return {
        "type": "paragraph",
        "style_name": style_name,
        "list_type": _resolve_list_type(style_name, p_pr, numbering),
        "runs": runs,
        "text": text,
    }


def _parse_table_block(
    element: ET.Element,
    styles: dict[str, str],
    numbering: dict[str, dict[str, str]],
) -> dict[str, Any]:
    rows: list[list[list[dict[str, Any]]]] = []
    for row_node in element.findall("w:tr", NS):
        row: list[list[dict[str, Any]]] = []
        for cell_node in row_node.findall("w:tc", NS):
            cell_blocks: list[dict[str, Any]] = []
            for child in cell_node:
                block = _parse_block(child, styles, numbering)
                if block:
                    cell_blocks.append(block)
            row.append(cell_blocks)
        rows.append(row)
    return {"type": "table", "rows": rows}


def _parse_runs_from_container(container: ET.Element) -> list[dict[str, Any]]:
    runs: list[dict[str, Any]] = []
    for child in container:
        if child.tag == W + "r":
            run = _parse_run(child)
            if run["text"]:
                runs.append(run)
        elif child.tag not in {W + "tbl"} and list(child):
            runs.extend(_parse_runs_from_container(child))
    return runs


def _parse_run(element: ET.Element) -> dict[str, Any]:
    run_props = element.find("w:rPr", NS)
    text_parts: list[str] = []
    for child in element:
        if child.tag == W + "t":
            text_parts.append(child.text or "")
        elif child.tag == W + "tab":
            text_parts.append("\t")
        elif child.tag in {W + "br", W + "cr"}:
            text_parts.append("\n")

    return {
        "text": "".join(text_parts),
        "bold": run_props is not None and run_props.find("w:b", NS) is not None,
        "italic": run_props is not None and run_props.find("w:i", NS) is not None,
        "underline": run_props is not None and run_props.find("w:u", NS) is not None,
    }


def _resolve_list_type(style_name: str, p_pr: ET.Element | None, numbering: dict[str, dict[str, str]]) -> str | None:
    normalized = (style_name or "").strip().lower()
    if "list bullet" in normalized or normalized.endswith("bullet") or "bullet" in normalized:
        return "ul"
    if "list number" in normalized or normalized.endswith("number") or "number" in normalized:
        return "ol"

    if p_pr is None:
        return None

    num_pr = p_pr.find("w:numPr", NS)
    if num_pr is None:
        return None

    num_id_node = num_pr.find("w:numId", NS)
    level_node = num_pr.find("w:ilvl", NS)
    num_id = num_id_node.get(W + "val") if num_id_node is not None else ""
    level = level_node.get(W + "val") if level_node is not None else "0"
    num_fmt = numbering.get(num_id, {}).get(level) or numbering.get(num_id, {}).get("0", "")

    if not num_fmt:
        return "ul"

    if num_fmt.lower() == "bullet":
        return "ul"
    return "ol"


def _get_block_tag(style_name: str) -> tuple[str, str | None]:
    normalized = (style_name or "").strip().lower()

    if normalized == "title":
        return "h1", "docx-preview-title"
    if normalized == "subtitle":
        return "p", "docx-preview-subtitle"
    if normalized.startswith("heading"):
        match = re.search(r"(\d+)", normalized)
        level = 1
        if match:
            level = min(max(int(match.group(1)), 1), 6)
        return f"h{level}", None

    return "p", None


def _render_blocks_to_html(blocks: list[dict[str, Any]]) -> str:
    parts: list[str] = []
    current_list_type: str | None = None

    def close_list() -> None:
        nonlocal current_list_type
        if current_list_type:
            parts.append(f"</{current_list_type}>")
            current_list_type = None

    for block in blocks:
        if block["type"] == "paragraph":
            text_html = _render_runs_to_html(block.get("runs", [])).strip()
            list_type = block.get("list_type")

            if list_type:
                if current_list_type != list_type:
                    close_list()
                    current_list_type = list_type
                    parts.append(f"<{list_type}>")
                if text_html:
                    parts.append(f"<li>{text_html}</li>")
                continue

            close_list()
            if not text_html:
                parts.append('<div class="docx-preview-spacer" aria-hidden="true"></div>')
                continue

            tag, css_class = _get_block_tag(block.get("style_name", ""))
            if css_class:
                parts.append(f'<{tag} class="{css_class}">{text_html}</{tag}>')
            else:
                parts.append(f"<{tag}>{text_html}</{tag}>")
            continue

        close_list()
        table_html = _render_parsed_table(block)
        if table_html:
            parts.append(table_html)

    close_list()

    if not parts:
        parts.append("<p>No previewable content found.</p>")
    return '<article class="docx-preview-fragment">' + "".join(parts) + "</article>"


def _render_runs_to_html(runs: list[dict[str, Any]]) -> str:
    chunks: list[str] = []
    for run in runs:
        text = _escape_run_text(run.get("text", ""))
        if not text:
            continue
        if run.get("bold"):
            text = f"<strong>{text}</strong>"
        if run.get("italic"):
            text = f"<em>{text}</em>"
        if run.get("underline"):
            text = f"<u>{text}</u>"
        chunks.append(text)
    return "".join(chunks)


def _escape_run_text(value: str) -> str:
    if not value:
        return ""
    escaped = html.escape(value)
    escaped = escaped.replace("\t", "&emsp;")
    return escaped.replace("\n", "<br>")


def _render_parsed_table(block: dict[str, Any]) -> str:
    rows = block.get("rows", [])
    if not rows:
        return ""

    rendered_rows = [[_render_parsed_table_cell(cell) for cell in row] for row in rows]
    if not any(any(cell.strip() for cell in row) for row in rendered_rows):
        return ""

    head_cells = "".join(f"<th>{cell or '&nbsp;'}</th>" for cell in rendered_rows[0])
    body_rows = rendered_rows[1:] if len(rendered_rows) > 1 else []
    body_html = "".join(
        "<tr>" + "".join(f"<td>{cell or '&nbsp;'}</td>" for cell in row) + "</tr>"
        for row in body_rows
    )

    parts = [
        '<div class="docx-preview-table-wrap"><table class="docx-preview-table">',
        f"<thead><tr>{head_cells}</tr></thead>",
    ]
    if body_html:
        parts.append(f"<tbody>{body_html}</tbody>")
    parts.append("</table></div>")
    return "".join(parts)


def _render_parsed_table_cell(blocks: list[dict[str, Any]]) -> str:
    chunks: list[str] = []
    for block in blocks:
        if block["type"] == "paragraph":
            html_text = _render_runs_to_html(block.get("runs", [])).strip()
            if html_text:
                chunks.append(html_text)
        else:
            nested = _render_parsed_table(block)
            if nested:
                chunks.append(nested)
    return "<br>".join(chunks)


def _blocks_to_text(blocks: list[dict[str, Any]]) -> str:
    lines: list[str] = []
    for block in blocks:
        if block["type"] == "paragraph":
            text = "".join(run.get("text", "") for run in block.get("runs", [])).strip()
            if not text:
                lines.append("")
                continue
            prefix = ""
            if block.get("list_type") == "ul":
                prefix = "- "
            elif block.get("list_type") == "ol":
                prefix = "1. "
            lines.append(prefix + text)
            continue

        for row in block.get("rows", []):
            row_values = [_cell_blocks_to_text(cell).strip() for cell in row]
            if any(value for value in row_values):
                lines.append(" | ".join(value for value in row_values if value))

    return _collapse_blank_lines(lines)


def _cell_blocks_to_text(blocks: list[dict[str, Any]]) -> str:
    parts: list[str] = []
    for block in blocks:
        if block["type"] == "paragraph":
            text = "".join(run.get("text", "") for run in block.get("runs", [])).strip()
            if text:
                parts.append(text)
        else:
            nested = _blocks_to_text([block]).strip()
            if nested:
                parts.append(nested)
    return "\n".join(parts)


def _collapse_blank_lines(lines: list[str]) -> str:
    normalized: list[str] = []
    blank_count = 0
    for line in lines:
        stripped = line.strip()
        if not stripped:
            blank_count += 1
            if blank_count > 1:
                continue
            normalized.append("")
            continue
        blank_count = 0
        normalized.append(stripped)
    return "\n".join(normalized).strip()


def _convert_via_docx2pdf(source: Path, target: Path) -> None:
    try:
        from docx2pdf import convert as convert_docx2pdf
    except ModuleNotFoundError as exc:
        raise RuntimeError("docx2pdf is not installed.") from exc

    temp_target = target.with_name(target.stem + ".build.pdf")
    temp_target.unlink(missing_ok=True)

    try:
        convert_docx2pdf(str(source), str(temp_target))
    except Exception as exc:
        temp_target.unlink(missing_ok=True)
        raise RuntimeError(f"docx2pdf failed: {exc}") from exc

    if not temp_target.exists():
        raise RuntimeError("docx2pdf did not produce a PDF file.")

    temp_target.replace(target)


def _convert_via_word_com(source: Path, target: Path) -> None:
    if os.name != "nt":
        raise RuntimeError("Word COM conversion is only available on Windows.")

    powershell = shutil.which("powershell.exe") or shutil.which("powershell")
    if not powershell:
        raise RuntimeError("PowerShell is not available for Word COM conversion.")

    script = r"""
param(
    [Parameter(Mandatory=$true)][string]$Source,
    [Parameter(Mandatory=$true)][string]$Target
)

$word = $null
$document = $null

try {
    $resolvedSource = [System.IO.Path]::GetFullPath($Source)
    $resolvedTarget = [System.IO.Path]::GetFullPath($Target)
    $targetDir = Split-Path -Parent $resolvedTarget
    if (-not (Test-Path $targetDir)) {
        New-Item -ItemType Directory -Path $targetDir | Out-Null
    }

    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $document = $word.Documents.Open($resolvedSource, $false, $true)
    $document.ExportAsFixedFormat($resolvedTarget, 17)
}
finally {
    if ($document -ne $null) {
        $document.Close($false)
    }
    if ($word -ne $null) {
        $word.Quit()
    }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
"""

    with tempfile.TemporaryDirectory(prefix="docx-word-pdf-") as temp_dir:
        script_path = Path(temp_dir) / "convert.ps1"
        script_path.write_text(script, encoding="utf-8")
        cmd = [
            powershell,
            "-NoProfile",
            "-NonInteractive",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script_path),
            str(source),
            str(target),
        ]
        try:
            subprocess.run(cmd, check=True, capture_output=True, text=True, timeout=180)
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError("Word COM conversion timed out.") from exc
        except subprocess.CalledProcessError as exc:
            stderr = (exc.stderr or "").strip()
            stdout = (exc.stdout or "").strip()
            detail = stderr or stdout or "Microsoft Word could not export the document."
            raise RuntimeError(f"Word COM conversion failed: {detail}") from exc

    if not target.exists():
        raise RuntimeError("Word COM did not produce a PDF file.")


def _convert_via_soffice(source: Path, target: Path) -> None:
    soffice = shutil.which("soffice")
    if not soffice:
        raise RuntimeError("LibreOffice soffice is not available.")

    with tempfile.TemporaryDirectory(prefix="docx-pdf-preview-") as temp_dir:
        temp_path = Path(temp_dir)
        cmd = [
            soffice,
            "--headless",
            "--convert-to",
            "pdf:writer_pdf_Export",
            "--outdir",
            str(temp_path),
            str(source),
        ]
        try:
            subprocess.run(cmd, check=True, capture_output=True, text=True, timeout=180)
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError("LibreOffice conversion timed out.") from exc
        except subprocess.CalledProcessError as exc:
            stderr = (exc.stderr or "").strip()
            raise RuntimeError(f"LibreOffice conversion failed: {stderr or exc}") from exc

        produced = temp_path / (source.stem + ".pdf")
        if not produced.exists():
            raise RuntimeError("LibreOffice did not produce a PDF file.")

        shutil.move(str(produced), str(target))


def _convert_via_browser_print(source: Path, target: Path) -> None:
    browser = _find_browser_executable()
    if not browser:
        raise RuntimeError("No headless browser was found for HTML-to-PDF fallback.")

    html_fragment = render_docx_to_html(str(source))
    page_html = PRINT_PAGE_HTML.format(title=html.escape(source.name), body=html_fragment)

    with tempfile.TemporaryDirectory(prefix="docx-browser-pdf-") as temp_dir:
        temp_path = Path(temp_dir)
        html_path = temp_path / "preview.html"
        output_path = temp_path / "preview.pdf"
        profile_dir = temp_path / "browser-profile"
        html_path.write_text(page_html, encoding="utf-8")
        file_url = html_path.resolve().as_uri()

        commands = [
            [
                str(browser),
                "--headless=new",
                "--disable-gpu",
                "--disable-dev-shm-usage",
                "--no-sandbox",
                "--run-all-compositor-stages-before-draw",
                "--virtual-time-budget=4000",
                f"--user-data-dir={profile_dir}",
                "--no-pdf-header-footer",
                f"--print-to-pdf={output_path}",
                file_url,
            ],
            [
                str(browser),
                "--headless",
                "--disable-gpu",
                "--disable-dev-shm-usage",
                "--no-sandbox",
                "--run-all-compositor-stages-before-draw",
                "--virtual-time-budget=4000",
                f"--user-data-dir={profile_dir}",
                "--print-to-pdf-no-header",
                f"--print-to-pdf={output_path}",
                file_url,
            ],
        ]

        errors: list[str] = []
        for cmd in commands:
            output_path.unlink(missing_ok=True)
            try:
                subprocess.run(cmd, check=True, capture_output=True, text=True, timeout=180)
            except subprocess.TimeoutExpired as exc:
                errors.append("browser print timed out")
                continue
            except subprocess.CalledProcessError as exc:
                stderr = (exc.stderr or "").strip()
                stdout = (exc.stdout or "").strip()
                errors.append(stderr or stdout or str(exc))
                continue

            if output_path.exists():
                shutil.move(str(output_path), str(target))
                return

        detail = " ; ".join(error for error in errors if error)
        raise RuntimeError(f"Headless browser PDF export failed. {detail}".strip())


def _find_browser_executable() -> Path | None:
    seen: set[str] = set()

    def candidates() -> Iterator[Path]:
        browser_names = (
            "msedge",
            "msedge.exe",
            "chrome",
            "chrome.exe",
            "google-chrome",
            "chromium",
            "chromium-browser",
        )
        for name in browser_names:
            resolved = shutil.which(name)
            if resolved:
                yield Path(resolved)

        windows_roots = [os.environ.get("PROGRAMFILES"), os.environ.get("PROGRAMFILES(X86)"), os.environ.get("LOCALAPPDATA")]
        for root in windows_roots:
            if not root:
                continue
            base = Path(root)
            for relative in (
                Path("Microsoft/Edge/Application/msedge.exe"),
                Path("Google/Chrome/Application/chrome.exe"),
                Path("Chromium/Application/chrome.exe"),
            ):
                candidate = base / relative
                if candidate.exists():
                    yield candidate

        for candidate in (
            Path("/usr/bin/google-chrome"),
            Path("/usr/bin/chromium"),
            Path("/usr/bin/chromium-browser"),
            Path("/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge"),
            Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"),
        ):
            if candidate.exists():
                yield candidate

    for candidate in candidates():
        key = str(candidate).lower()
        if key in seen:
            continue
        seen.add(key)
        return candidate
    return None
